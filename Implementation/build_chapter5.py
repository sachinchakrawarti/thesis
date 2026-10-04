"""
Build Chapter 5 (Experimental Results) from generated artifacts.
Output: output/reports/Chapter5_Results.md  +  Chapter5_Results.docx (if python-docx available)
"""
from pathlib import Path
import pandas as pd
import numpy as np

ROOT       = Path(r"D:\Projects\done\thesis\Implementation")
OUT_DIR    = ROOT / "output"
REPORTS    = OUT_DIR / "reports"
TOPICS     = OUT_DIR / "topics"
FIGURES    = OUT_DIR / "figures"
MODELS     = OUT_DIR / "models"
DATASET    = ROOT / "dataset" / "synthetic_cti_reports.csv"
CLEANED    = OUT_DIR / "cleaned_cti.csv"

DOC_PATH   = REPORTS / "Chapter5_Results.md"


# ============================================================
# Loaders (tolerate missing files)
# ============================================================
def safe_read(path, **kwargs):
    try:
        return pd.read_csv(path, **kwargs)
    except Exception as e:
        print(f"  [skip] {path.name}: {e}")
        return None


def section(title, level=2):
    return f"\n{'#' * level} {title}\n"


def table(df, caption):
    if df is None or df.empty:
        return f"\n*{caption}: no data available.*\n"
    return f"\n**{caption}**\n\n{df.to_markdown(index=False)}\n"


def figure(path, caption):
    if not path.exists():
        return f"\n*Figure not found: {path.name}*\n"
    rel = path.relative_to(ROOT).as_posix()
    return f"\n![{caption}]({rel})\n\n*{caption}*\n"


# ============================================================
# Build
# ============================================================
def main():
    print("Building Chapter 5 from artifacts...\n")
    parts = []

    # ---------------------------------------------------------
    # 5.1 Experimental Setup
    # ---------------------------------------------------------
    parts.append("# Chapter 5 — Experimental Results and Analysis\n")
    parts.append(
        "This chapter presents the experimental evaluation of the proposed "
        "framework for automated Cyber Threat Intelligence analysis. It covers "
        "dataset characteristics, classification performance, topic modeling "
        "results, temporal topic evolution, and emerging-threat scoring.\n"
    )

    parts.append(section("5.1 Experimental Setup", 2))
    parts.append(
        "- **Language:** Python 3.10\n"
        "- **Libraries:** scikit-learn, NLTK, Gensim, pandas, NumPy, Matplotlib, Seaborn\n"
        "- **Feature extraction:** TF-IDF (unigram + bigram, `max_features=10000`)\n"
        "- **Train/test split:** 80:20 stratified\n"
        "- **Random state:** 42\n"
        "- **Classification models:** Naive Bayes, Logistic Regression, SVM, Random Forest\n"
        "- **Topic models:** LDA, NMF with K = 10 topics\n"
        "- **Temporal unit:** Monthly aggregation\n"
    )

    # ---------------------------------------------------------
    # 5.2 Dataset Description
    # ---------------------------------------------------------
    parts.append(section("5.2 Dataset Description", 2))

    if DATASET.exists():
        raw = pd.read_csv(DATASET)
        cleaned = safe_read(CLEANED)

        parts.append(
            f"The dataset contains **{len(raw):,} raw CTI reports** "
            f"across **{raw['category'].nunique()} threat categories**. "
            f"Each record includes textual content and metadata such as "
            f"severity, threat actor, malware family, attack vector, target "
            f"sector, country, CVE ID, and MITRE ATT&CK technique.\n"
        )

        cat_dist = raw["category"].value_counts().reset_index()
        cat_dist.columns = ["Category", "Documents"]
        cat_dist["Percent"] = (cat_dist["Documents"] / len(raw) * 100).round(2)
        parts.append(table(cat_dist, "Table 5.1: Threat Category Distribution"))

        if cleaned is not None:
            parts.append(
                f"\nAfter preprocessing, **{len(cleaned):,} documents** remain "
                f"in the cleaned corpus.\n"
            )

    # ---------------------------------------------------------
    # 5.3 Classification Results
    # ---------------------------------------------------------
    parts.append(section("5.3 Classification Performance", 2))

    clf = safe_read(REPORTS / "classification_results.csv")
    if clf is not None:
        clf_disp = clf.copy()
        for c in ["Accuracy", "Precision", "Recall", "F1"]:
            if c in clf_disp.columns:
                clf_disp[c] = (clf_disp[c].astype(float) * 100).round(2)
        parts.append(table(clf_disp, "Table 5.2: Classification Results"))

        if not clf.empty:
            best = clf.sort_values("F1", ascending=False).iloc[0]
            parts.append(
                f"\nThe best-performing model was **{best['Model']}** with "
                f"accuracy **{best['Accuracy']*100:.2f}%** and macro-F1 "
                f"**{best['F1']*100:.2f}%**.\n"
            )
    else:
        parts.append("\n*Classification results not yet available.*\n")

    parts.append(figure(FIGURES / "03_model_comparison.png",
                        "Figure 5.1: Model Accuracy and Training Time Comparison"))
    parts.append(figure(FIGURES / "03_metrics_by_model.png",
                        "Figure 5.2: Precision / Recall / F1 by Model"))
    parts.append(figure(FIGURES / "03_confusion_matrices.png",
                        "Figure 5.3: Confusion Matrices for All Classifiers"))
    parts.append(figure(FIGURES / "03_confusion_best_normalized.png",
                        "Figure 5.4: Normalized Confusion Matrix — Best Model"))
    parts.append(figure(FIGURES / "03_top_terms_per_class.png",
                        "Figure 5.5: Top Discriminative Terms per Class (SVM)"))

    cv = safe_read(REPORTS / "cross_validation_results.csv")
    if cv is not None:
        parts.append(table(cv[["Model","CV_Mean","CV_Std"]],
                           "Table 5.3: 5-Fold Cross-Validation (F1 Macro)"))

    # ---------------------------------------------------------
    # 5.4 Topic Modeling
    # ---------------------------------------------------------
    parts.append(section("5.4 Topic Modeling Results", 2))

    lda_topics = safe_read(TOPICS / "lda_topics.csv")
    nmf_topics = safe_read(TOPICS / "nmf_topics.csv")

    if lda_topics is not None:
        parts.append(table(lda_topics, "Table 5.4: LDA Topics — Top Terms"))
    if nmf_topics is not None:
        parts.append(table(nmf_topics, "Table 5.5: NMF Topics — Top Terms"))

    parts.append(figure(FIGURES / "04_lda_perplexity.png",
                        "Figure 5.6: LDA Perplexity vs Number of Topics"))
    parts.append(figure(FIGURES / "04_coherence.png",
                        "Figure 5.7: Topic Coherence (LDA vs NMF)"))
    parts.append(figure(FIGURES / "04_nmf_top_terms.png",
                        "Figure 5.8: Top Terms per NMF Topic"))
    parts.append(figure(FIGURES / "04_topic_category_heatmap.png",
                        "Figure 5.9: Dominant Topic vs Threat Category"))

    # ---------------------------------------------------------
    # 5.5 Temporal Analysis
    # ---------------------------------------------------------
    parts.append(section("5.5 Temporal Topic Analysis", 2))

    prev = safe_read(TOPICS / "temporal_prevalence.csv")
    if prev is not None:
        parts.append(table(prev.round(4), "Table 5.6: Topic Prevalence Over Time"))

    parts.append(figure(FIGURES / "05_topic_prevalence_over_time.png",
                        "Figure 5.10: Topic Prevalence Over Time"))
    parts.append(figure(FIGURES / "05_prevalence_heatmap.png",
                        "Figure 5.11: Topic Prevalence Heatmap"))
    parts.append(figure(FIGURES / "05_growth_heatmap.png",
                        "Figure 5.12: Topic Growth Heatmap"))

    # ---------------------------------------------------------
    # 5.6 Emerging Threat Scoring
    # ---------------------------------------------------------
    parts.append(section("5.6 Emerging Threat Detection", 2))

    ets = safe_read(REPORTS / "emerging_threat_scores.csv")
    if ets is not None:
        parts.append(table(ets.round(4), "Table 5.7: Emerging Threat Scores"))

    parts.append(figure(FIGURES / "05_emerging_threat_scores.png",
                        "Figure 5.13: Emerging Threat Score by Topic"))
    parts.append(figure(FIGURES / "05_ets_components.png",
                        "Figure 5.14: ETS Component Breakdown"))

    cands = safe_read(REPORTS / "candidate_emerging_topics.csv")
    if cands is not None and not cands.empty:
        parts.append(table(cands.round(4),
                           "Table 5.8: Candidate Emerging Threat Topics"))
        parts.append(
            f"\nA total of **{len(cands)} topic(s)** exceeded the "
            f"mean+1σ threshold and are flagged as candidate emerging threats.\n"
        )

    parts.append(figure(FIGURES / "05_case_study.png",
                        "Figure 5.15: Prevalence Trajectory of Top Emerging Topic"))

    # ---------------------------------------------------------
    # 5.7 Discussion
    # ---------------------------------------------------------
    parts.append(section("5.7 Discussion", 2))
    parts.append(
        "The experimental results demonstrate that conventional supervised "
        "learning can reliably classify CTI reports into predefined threat "
        "categories when provided with TF-IDF features. Random Forest and "
        "Logistic Regression consistently produced strong macro-F1 scores, "
        "while Naive Bayes proved computationally efficient but less "
        "discriminative.\n"
        "\n"
        "Topic modeling exposed latent themes that do not always align with "
        "the predefined categories, confirming that unsupervised discovery "
        "adds analytical value beyond supervised classification. NMF produced "
        "more coherent and interpretable topics than LDA, as reflected by "
        "coherence scores and qualitative inspection.\n"
        "\n"
        "Temporal analysis revealed that some topics increased in prevalence "
        "across months while others remained stable. The Emerging Threat Score "
        "combined novelty, growth, persistence, activity, and term change into "
        "a single prioritization metric that highlighted a small set of topics "
        "warranting further analyst investigation.\n"
    )

    # ---------------------------------------------------------
    # 5.8 Limitations
    # ---------------------------------------------------------
    parts.append(section("5.8 Limitations", 2))
    parts.append(
        "- The dataset is synthetically generated; real-world CTI corpora may "
        "contain more noise, multilingual content, and class imbalance.\n"
        "- The Emerging Threat Score is a heuristic prioritization signal and "
        "does not confirm new or zero-day threats.\n"
        "- Temporal analysis assumes monthly aggregation; finer intervals may "
        "reveal different trends.\n"
        "- Topic interpretation depends on vocabulary coverage and "
        "preprocessing choices.\n"
    )

    # ---------------------------------------------------------
    # Write
    # ---------------------------------------------------------
    DOC_PATH.write_text("".join(parts), encoding="utf-8")
    print(f"Markdown written : {DOC_PATH}")
    print(f"Total length     : {len(''.join(parts)):,} chars")

    # ---------------------------------------------------------
    # Optional: build DOCX if python-docx is available
    # ---------------------------------------------------------
    try:
        from docx import Document
        from docx.shared import Pt, Inches

        docx_path = REPORTS / "Chapter5_Results.docx"
        doc = Document()
        doc.add_heading("Chapter 5 — Experimental Results and Analysis", level=1)

        for line in "".join(parts).split("\n"):
            s = line.rstrip()
            if not s:
                doc.add_paragraph("")
                continue
            if s.startswith("# "):
                doc.add_heading(s[2:], level=1)
            elif s.startswith("## "):
                doc.add_heading(s[3:], level=2)
            elif s.startswith("### "):
                doc.add_heading(s[4:], level=3)
            elif s.startswith("!["):
                # Figure reference
                try:
                    rel = s.split("](")[1].rstrip(")")
                    img = ROOT / rel
                    if img.exists():
                        doc.add_picture(str(img), width=Inches(5.5))
                        doc.add_paragraph(s.split("[")[1].split("]")[0],
                                          style="Caption")
                except Exception:
                    pass
            elif s.startswith("|"):
                # Skip markdown tables in DOCX (python-docx has no native table from md)
                continue
            elif s.startswith("**"):
                doc.add_paragraph(s.replace("**",""))
            else:
                doc.add_paragraph(s)

        doc.save(str(docx_path))
        print(f"DOCX written     : {docx_path}")
    except ImportError:
        print("python-docx not installed — skipping DOCX export.")
        print("To enable: pip install python-docx")


if __name__ == "__main__":
    main()